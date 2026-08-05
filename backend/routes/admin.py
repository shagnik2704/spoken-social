# ============================================================
# Admin routes — platform-wide aggregations + feedback analytics
# ============================================================
# 1. /platform-summary  — Overview + Analytics tab data
# 2. /feedback-summary   — Aggregated student feedback
# 3. /students/{id}/feedback — Per-student feedback detail
# ============================================================

from cache import get_cached,set_cached
from fastapi import APIRouter, Depends, Query
from typing import Optional
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, case, cast, Date
from datetime import datetime, date
import calendar
import openpyxl

from excel_export import workbook_response, write_styled_sheet
from dashboard_common import (
    DEFAULT_TUTORIAL_COUNT,
    PROGRESS_SERIES,
    STATUS_SERIES,
    chart_series,
    course_xp_sum,
    initials as build_initials,
    progress_bucket,
)
from database import get_db
from demo_data import get_mock_comparison_data, generate_deterministic_mock_metrics
from services.feedback_summary import build_feedback_summary
from models import User, College, UserTutorial, Tutorial, UserCourse, Programme, Course, Batch, ProgrammeCourse
from models.feedback import TutorialFeedback
from models.xp_event import XPEvent
from dependencies import require_admin

router = APIRouter(prefix="/api/admin", tags=["admin"])


@router.get("/platform-summary")
async def get_platform_summary(
    course_id: str = Query("linux"),
    batch_id: Optional[str] = Query(None),
    state: Optional[str] = Query(None),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Returns a single, pre-computed payload for the admin Overview + Analytics tabs.
    Now course and batch aware.
    """

    cache_key = f"admin:platform_summary:{course_id}:{batch_id or 'all'}:{state or 'all'}"
    cached = await get_cached(cache_key)
    if cached:
        return cached

    tot_tut_res = await db.execute(select(func.count(Tutorial.id)).filter(Tutorial.course_id == course_id))
    total_tutorials = tot_tut_res.scalar() or DEFAULT_TUTORIAL_COUNT

    # ----- Per-student metrics (one query) -----
    query = select(
        User.id,
        User.display_name,
        User.username,
        course_xp_sum().label("xp_total"),
        User.college_id,
        func.count(
            case((UserTutorial.watched == True, Tutorial.id), else_=None)
        ).label("watched"),
        func.count(
            case((UserTutorial.practiced == True, Tutorial.id), else_=None)
        ).label("completed"),
        func.count(
            case((UserTutorial.minitask_opted == True, Tutorial.id), else_=None)
        ).label("minitasks_opted"),
    ).outerjoin(
        UserTutorial, UserTutorial.user_id == User.id
    ).outerjoin(
        Tutorial,
        (Tutorial.id == UserTutorial.tutorial_id) & (Tutorial.course_id == course_id)
    ).join(
        UserCourse, UserCourse.user_id == User.id
    ).filter(
        UserCourse.course_id == course_id
    )

    if batch_id:
        query = query.filter(UserCourse.batch_id == batch_id)

    if state:
        query = query.join(College, College.id == User.college_id).filter(
            User.role == "student",
            User.is_test_account != True,
            College.location.ilike(f"%{state}%")
        )
    else:
        query = query.filter(User.role == "student", User.is_test_account != True)

    query = query.group_by(User.id)
    res_metrics = await db.execute(query)
    student_metrics = res_metrics.all()

    # ----- Build college name lookup -----
    college_query = (
        select(College.id, College.name)
        .join(User, User.college_id == College.id)
        .join(UserCourse, UserCourse.user_id == User.id)
        .filter(
            User.role == "student",
            User.is_test_account != True,
            UserCourse.course_id == course_id
        )
    )
    if batch_id:
        college_query = college_query.filter(UserCourse.batch_id == batch_id)
    if state:
        college_query = college_query.filter(College.location.ilike(f"%{state}%"))
        
    colleges_res = await db.execute(college_query.distinct())
    colleges_loaded = colleges_res.all()
    college_names = {c.id: c.name for c in colleges_loaded}

    # ----- Compute derived metrics -----
    students_enriched = []
    status_counts = {"completed": 0, "active": 0, "at_risk": 0, "inactive": 0, "not_started": 0}
    progress_buckets = [0] * len(PROGRESS_SERIES)
    total_progress = 0

    for row in student_metrics:
        uid, name, username, xp, college_id, watched, completed, minitasks_opted = row
        progress = round((completed / total_tutorials) * 100) if total_tutorials > 0 else 0
        total_progress += progress

        # Derive status
        if completed >= total_tutorials:
            status = "completed"
        elif watched > total_tutorials * 0.3:
            status = "active"
        elif watched > 0:
            status = "at_risk"
        else:
            status = "inactive" if xp > 0 else "not_started"

        status_counts[status] += 1
        progress_buckets[progress_bucket(progress)] += 1

        initials = build_initials(name, username)

        students_enriched.append({
            "id": str(uid),
            "name": name or username or "",
            "initials": initials,
            "xp": xp,
            "progress": progress,
            "status": status,
            "watched": watched,
            "completed": completed,
            "minitasksOpted": minitasks_opted,
            "collegeName": college_names.get(college_id, "Unknown"),
        })

    total_students = len(students_enriched)
    avg_progress = round(total_progress / total_students) if total_students > 0 else 0
    avg_xp = round(sum(s["xp"] for s in students_enriched) / total_students) if total_students > 0 else 0
    completion_rate = round((status_counts["completed"] / total_students) * 100) if total_students > 0 else 0

    # ----- Top 10 performers by XP (only those with XP > 0) -----
    top_performers_eligible = [s for s in students_enriched if s["xp"] > 0]
    top_performers = sorted(top_performers_eligible, key=lambda s: s["xp"], reverse=True)[:10]

    # ----- At-risk students (at_risk + inactive + not_started, sorted by lowest progress) -----
    has_any_activity = any(s["xp"] > 0 or s["watched"] > 0 for s in students_enriched)
    allowed_statuses = ("at_risk", "inactive", "not_started") if has_any_activity else ("at_risk", "inactive")
    at_risk = sorted(
        [s for s in students_enriched if s["status"] in allowed_statuses],
        key=lambda s: s["progress"],
    )[:10]

    # ----- Tutorial-wise completion rates -----
    # The numerator must count only students in the SAME population as total_students
    # (course + batch + role + non-test + optional state). Previously the numerator
    # ignored batch_id/role/test filters, so `rate` mixed a batch-filtered denominator
    # with an unfiltered numerator. Build the population once and scope the count to it.
    pop_query = (
        select(User.id)
        .join(UserCourse, UserCourse.user_id == User.id)
        .filter(
            User.role == "student",
            User.is_test_account != True,
            UserCourse.course_id == course_id,
        )
    )
    if batch_id:
        pop_query = pop_query.filter(UserCourse.batch_id == batch_id)
    if state:
        pop_query = pop_query.join(College, College.id == User.college_id).filter(
            College.location.ilike(f"%{state}%")
        )

    stats_query = (
        select(
            Tutorial.tutorial_number,
            Tutorial.title,
            func.count(
                case((UserTutorial.practiced == True, UserTutorial.user_id), else_=None)
            ).label("completed_count"),
        )
        .select_from(Tutorial)
        .outerjoin(
            UserTutorial,
            (UserTutorial.tutorial_id == Tutorial.id)
            & (UserTutorial.user_id.in_(pop_query.distinct())),
        )
        .filter(Tutorial.course_id == course_id)
        .group_by(Tutorial.tutorial_number, Tutorial.title)
        .order_by(Tutorial.tutorial_number)
    )
    res_stats = await db.execute(stats_query)
    tutorial_stats = res_stats.all()

    tutorial_completion = [
        {
            "number": t.tutorial_number,
            "name": t.title,
            "completedCount": t.completed_count,
            "rate": round((t.completed_count / total_students) * 100) if total_students > 0 else 0,
        }
        for t in tutorial_stats
    ]

    result = {
        "metrics": {
            "totalStudents": total_students,
            "totalColleges": len(colleges_loaded),
            "avgProgress": avg_progress,
            "avgXP": avg_xp,
            "completionRate": completion_rate,
            "totalCompleted": status_counts["completed"],
            "totalActive": status_counts["active"],
            "totalAtRisk": status_counts["at_risk"],
            "totalInactive": status_counts["inactive"] + status_counts["not_started"],
            # Broken out so the dashboard can tell "behind" apart from "never began" — a
            # cohort where nobody has started needs a different message, not a count of
            # students needing attention.
            "totalNotStarted": status_counts["not_started"],
        },
        "statusDistribution": chart_series(STATUS_SERIES, [
            status_counts["completed"],
            status_counts["active"],
            status_counts["at_risk"],
            status_counts["inactive"] + status_counts["not_started"],
        ]),
        "progressDistribution": chart_series(PROGRESS_SERIES, progress_buckets),
        "topPerformers": top_performers,
        "atRiskStudents": at_risk,
        "tutorialCompletion": tutorial_completion,
    }
    await set_cached(cache_key, result)
    return result


@router.get("/minitask-analysis")
async def get_minitask_analysis(
    course_id: str = Query("linux"),
    batch_id: Optional[str] = Query(None),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Returns aggregated minitask opt-in data by college and tutorial.
    """
    cache_key = f"admin:minitask:{course_id}:{batch_id or 'all'}"
    cached = await get_cached(cache_key)
    if cached is not None:
        return cached

    # Query all minitask opt-ins
    query = (
        select(
            College.name.label("college_name"),
            Tutorial.tutorial_number,
            func.count(UserTutorial.user_id).label("count")
        )
        .join(User, User.id == UserTutorial.user_id)
        .join(Tutorial, Tutorial.id == UserTutorial.tutorial_id)
        .join(UserCourse, UserCourse.user_id == User.id)
        .outerjoin(College, College.id == User.college_id)
        .filter(
            User.role == "student",
            User.is_test_account != True,
            UserTutorial.minitask_opted == True,
            Tutorial.course_id == course_id,
            UserCourse.course_id == course_id
        )
    )

    if batch_id:
        query = query.filter(UserCourse.batch_id == batch_id)
    
    query = query.group_by(College.name, Tutorial.tutorial_number).order_by(College.name, Tutorial.tutorial_number)
    
    res = await db.execute(query)
    rows = res.all()

    # Structure the data for frontend: {college_name: {tut_num: count}}
    analysis = {}
    for row in rows:
        cname = row.college_name or "Unknown"
        if cname not in analysis:
            analysis[cname] = {}
        analysis[cname][row.tutorial_number] = row.count

    await set_cached(cache_key, analysis, ttl=300)
    return analysis




@router.get("/colleges/export")
async def export_all_colleges(
    course_id: str = Query("linux"),
    batch_id: Optional[str] = Query(None),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Exports the macro view of all colleges and their stats to an Excel file.
    """
    # Import the same logic used to calculate their summary
    from routes.coordinator import _build_college_summary
    
    colleges_summary = await _build_college_summary(db, course_id=course_id, batch_id=batch_id)

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Colleges Report"

    write_styled_sheet(
        sheet,
        columns=[
            ("S.No.", 10),
            ("College Name", 45),
            ("Total Students Enrolled", 28),
            ("Active + Completed", 24),
            ("Fully Completed Students", 28),
            ("Students Opted for Minitasks", 30),
            ("Average Progress Percentage", 32),
        ],
        rows=[
            [
                index,
                col["name"],
                col["totalCourseStudents"],
                col["activeStudents"],
                col["completedStudents"],
                col.get("studentsWithMinitasks", 0),
                f"{col['avgProgress']}%",
            ]
            for index, col in enumerate(colleges_summary, start=1)
        ],
    )

    return workbook_response(wb, "Platform_Colleges_Report.xlsx")


# ============================================================
# Mini-Task Opt-in Export
# ============================================================

@router.get("/minitask-students/export")
async def export_minitask_students(
    course_id: str = Query("linux"),
    batch_id: Optional[str] = Query(None),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Exports a flat list of all students who opted for mini-tasks.
    One row per (student, tutorial) pair where minitask_opted = True.
    Columns: S.No., Student Name, Email, College, Tutorial #, Tutorial Title
    Sorted by tutorial number, then student name.

    Population MUST match the on-screen MiniTask analysis (get_minitask_analysis):
    enrolled students only (UserCourse join) and the same batch filter. Previously the
    export skipped both, so it included unenrolled/other-batch rows the screen never showed.
    """
    # Query all minitask opt-ins with student + tutorial + college info
    query = (
        select(
            User.display_name,
            User.username,
            User.email,
            College.name.label("college_name"),
            Tutorial.tutorial_number,
            Tutorial.title.label("tutorial_title"),
        )
        .join(UserTutorial, UserTutorial.user_id == User.id)
        .join(Tutorial, Tutorial.id == UserTutorial.tutorial_id)
        .join(UserCourse, UserCourse.user_id == User.id)
        .outerjoin(College, College.id == User.college_id)
        .filter(
            User.role == "student",
            User.is_test_account != True,
            UserTutorial.minitask_opted == True,
            Tutorial.course_id == course_id,
            UserCourse.course_id == course_id,
        )
    )
    if batch_id:
        query = query.filter(UserCourse.batch_id == batch_id)

    res = await db.execute(query.order_by(Tutorial.tutorial_number, User.display_name))
    rows = res.all()

    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "Mini-Task Opt-ins"

    # Left-align the text columns (name, email, college, tutorial title); centre the rest.
    TEXT_COLS = (2, 3, 4, 6)

    write_styled_sheet(
        sheet,
        columns=[
            ("S.No.", 10),
            ("Student Name", 35),
            ("Email", 40),
            ("College", 45),
            ("Tutorial #", 14),
            ("Tutorial Title", 45),
        ],
        rows=[
            [
                index,
                row.display_name or row.username or "",
                row.email or "",
                row.college_name or "Unknown",
                row.tutorial_number,
                row.tutorial_title,
            ]
            for index, row in enumerate(rows, start=1)
        ],
        left_align_cols=TEXT_COLS,
        autofilter=True,
        freeze_header=True,
    )

    # ── Sheet 2: Unique Students (one row per student) ──
    from collections import OrderedDict
    student_map = OrderedDict()
    for row in rows:
        email = row.email or ""
        if email not in student_map:
            student_map[email] = {
                "name": row.display_name or row.username or "",
                "email": email,
                "college": row.college_name or "Unknown",
                "tutorials": [],
            }
        student_map[email]["tutorials"].append(
            f"{row.tutorial_number}. {row.tutorial_title}"
        )

    write_styled_sheet(
        wb.create_sheet(title="Unique Students"),
        columns=[
            ("S.No.", 10),
            ("Student Name", 35),
            ("Email", 40),
            ("College", 45),
            ("Total Mini-Tasks", 18),
            ("Tutorials Opted", 70),
        ],
        rows=[
            [
                index,
                info["name"],
                info["email"],
                info["college"],
                len(info["tutorials"]),
                ", ".join(info["tutorials"]),
            ]
            for index, info in enumerate(student_map.values(), start=1)
        ],
        left_align_cols=TEXT_COLS,
        autofilter=True,
        freeze_header=True,
    )

    return workbook_response(wb, "MiniTask_OptIn_Report.xlsx")


# ============================================================
# College Daily Average XP (Line Chart Data)
# ============================================================

COLLEGE_COLORS = [
    "#8b5cf6", "#3b82f6", "#10b981", "#f59e0b", "#ef4444",
    "#ec4899", "#06b6d4", "#84cc16", "#f97316", "#6366f1",
    "#14b8a6", "#e11d48", "#0ea5e9", "#a855f7", "#d946ef",
]


@router.get("/college-daily-xp")
async def get_college_daily_xp(
    course_id: str = Query("linux"),
    month: Optional[str] = Query(None, description="Month in YYYY-MM format, defaults to current"),
    college_ids: Optional[str] = Query(None, description="Comma-separated college IDs"),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Returns daily average XP per college for a given month.
    For each day: total XP earned that day in this course, divided by the college's
    enrolled students — students who earned nothing count as zero. Dividing by earners
    instead would make one active student out of 300 chart the same as 300 out of 300.
    """
    # Parse month
    if month:
        try:
            year, mo = int(month.split("-")[0]), int(month.split("-")[1])
        except (ValueError, IndexError):
            year, mo = datetime.now().year, datetime.now().month
    else:
        year, mo = datetime.now().year, datetime.now().month

    days_in_month = calendar.monthrange(year, mo)[1]
    month_start = date(year, mo, 1)
    if mo == 12:
        month_end = date(year + 1, 1, 1)
    else:
        month_end = date(year, mo + 1, 1)

    # Cache only the default (all-colleges) chart — the per-selection variants are rare,
    # and the frontend now filters colleges client-side against this default.
    use_cache = not college_ids
    cache_key = f"admin:dailyxp:{course_id}:{year:04d}-{mo:02d}"
    if use_cache:
        cached = await get_cached(cache_key)
        if cached is not None:
            return cached

    # Get colleges
    college_filter_ids = None
    if college_ids:
        college_filter_ids = [cid.strip() for cid in college_ids.split(",") if cid.strip()]

    if college_filter_ids:
        colleges_query = select(College.id, College.name).filter(College.id.in_(college_filter_ids))
        res_colleges = await db.execute(colleges_query.order_by(College.name))
        colleges = res_colleges.all()
    else:
        # Limit default chart to top 10 active colleges by student enrollment count in this course
        college_query = (
            select(College.id, College.name)
            .join(User, User.college_id == College.id)
            .join(UserCourse, UserCourse.user_id == User.id)
            .filter(
                User.role == "student",
                User.is_test_account != True,
                UserCourse.course_id == course_id
            )
            .group_by(College.id, College.name)
            .order_by(func.count(User.id).desc())
            .limit(10)
        )
        res_colleges = await db.execute(college_query)
        # Sort them alphabetically for a cleaner UI presentation
        colleges = sorted(res_colleges.all(), key=lambda c: c.name)

    if not colleges:
        return {
            "month": f"{year:04d}-{mo:02d}",
            "daysInMonth": days_in_month,
            "colleges": [],
            "series": [{"day": d} for d in range(1, days_in_month + 1)],
        }

    college_list = []
    for i, c in enumerate(colleges):
        college_list.append({
            "id": c.id,
            "name": c.name,
            "color": COLLEGE_COLORS[i % len(COLLEGE_COLORS)],
        })

    # Subquery: per-student daily XP sum. Bucket by IST calendar day (created_at is
    # timestamptz; the audience is IST, so a "day" must be local, not UTC).
    from config import APP_TIMEZONE
    ist_date = cast(func.timezone(APP_TIMEZONE, XPEvent.created_at), Date)
    daily_student_xp = (
        select(
            User.college_id.label("college_id"),
            ist_date.label("xp_date"),
            XPEvent.user_id.label("user_id"),
            func.sum(XPEvent.xp_delta).label("daily_xp"),
        )
        .join(User, XPEvent.user_id == User.id)
        .join(UserCourse, UserCourse.user_id == User.id)
        .filter(
            User.role == "student",
            User.is_test_account != True,
            UserCourse.course_id == course_id,
            # Course-scope the POINTS too (not just the students): only sum XP earned in
            # this course. Events pre-dating course tagging are NULL and excluded.
            XPEvent.course_id == course_id,
            User.college_id.in_([c.id for c in colleges]),
            ist_date >= month_start,
            ist_date < month_end,
        )
        .group_by(User.college_id, ist_date, XPEvent.user_id)
        .subquery()
    )

    # Enrolled-student count per college — the denominator below. Same population as the
    # numerator's students (role, test-account and course filters all match).
    enrolled_res = await db.execute(
        select(User.college_id, func.count(func.distinct(User.id)).label("enrolled"))
        .join(UserCourse, UserCourse.user_id == User.id)
        .filter(
            User.role == "student",
            User.is_test_account != True,
            UserCourse.course_id == course_id,
            User.college_id.in_([c.id for c in colleges]),
        )
        .group_by(User.college_id)
    )
    enrolled_by_college = {row.college_id: row.enrolled for row in enrolled_res.all()}

    # Daily XP totals per college. Averaging happens in Python against the enrolled count
    # rather than as SQL AVG over the subquery, because AVG would divide by the number of
    # students who happened to earn that day — making one active student out of 300 chart
    # identically to 300 out of 300.
    result = await db.execute(
        select(
            daily_student_xp.c.college_id,
            daily_student_xp.c.xp_date,
            func.sum(daily_student_xp.c.daily_xp).label("total_xp"),
        )
        .group_by(daily_student_xp.c.college_id, daily_student_xp.c.xp_date)
        .order_by(daily_student_xp.c.xp_date)
    )
    rows = result.all()

    # Build series: one entry per day, with each college as a key
    series = []
    for day_num in range(1, days_in_month + 1):
        entry: dict = {"day": day_num}
        for c in colleges:
            entry[c.id] = 0
        series.append(entry)

    for row in rows:
        day_num = row.xp_date.day
        college_id = row.college_id
        enrolled = enrolled_by_college.get(college_id, 0)
        avg_xp = round(float(row.total_xp or 0) / enrolled, 1) if enrolled else 0
        if 1 <= day_num <= days_in_month:
            series[day_num - 1][college_id] = avg_xp

    result = {
        "month": f"{year:04d}-{mo:02d}",
        "daysInMonth": days_in_month,
        "colleges": college_list,
        "series": series,
    }
    if use_cache:
        await set_cached(cache_key, result, ttl=300)
    return result


# ============================================================
# Feedback Analytics (Filterable)
# ============================================================

@router.get("/feedback-summary")
async def get_feedback_summary(
    course_id: str = Query("linux"),
    tutorial_id: Optional[str] = Query(None, description="Filter by tutorial UUID"),
    step_type: Optional[str] = Query(None, description="Filter by step type"),
    reason: Optional[str] = Query(None, description="Filter by reason key"),
    college_id: Optional[str] = Query(None, description="Filter by college"),
    batch_id: Optional[str] = Query(None, description="Filter by batch"),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Filterable feedback analytics for the admin panel.
    Returns:
    1. Filter options (tutorials, step types, reasons) for dropdowns
    2. Aggregated counts based on current filters
    3. Filtered entries grouped by tutorial, with student details
    """
    cache_key = (
        f"fb:admin:{course_id}:{batch_id or '-'}:{college_id or '-'}:"
        f"{tutorial_id or '-'}:{step_type or '-'}:{reason or '-'}"
    )
    cached = await get_cached(cache_key)
    if cached is not None:
        return cached

    result = await build_feedback_summary(
        db,
        course_id=course_id,
        college_id=college_id,
        batch_id=batch_id,
        tutorial_id=tutorial_id,
        step_type=step_type,
        reason=reason,
    )
    await set_cached(cache_key, result, ttl=120)
    return result


@router.get("/students/{student_id}/feedback")
async def get_student_feedback(
    student_id: str,
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Returns detailed feedback data for a single student:
    1. Student info (name, college, roll number)
    2. Summary stats (total issues, most common reason, most problematic step)
    3. Reason breakdown with counts
    4. All entries chronologically with tutorial details
    """
    import uuid as uuid_mod
    from fastapi import HTTPException

    # A malformed id is a client error (400), not an uncaught ValueError → 500.
    try:
        sid = uuid_mod.UUID(student_id)
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid student ID format")

    # Get the student
    res_st = await db.execute(select(User).filter_by(id=sid))
    student = res_st.scalars().first()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    # Get college name
    college_name = "Unknown"
    if student.college_id:
        res_c = await db.execute(select(College.name).filter_by(id=student.college_id))
        college = res_c.first()
        if college:
            college_name = college.name

    initials = build_initials(student.display_name, student.username)

    student_info = {
        "id": str(student.id),
        "name": student.display_name or student.username,
        "initials": initials,
        "rollNo": student.username or "",
        "collegeName": college_name,
    }

    # Fetch all feedback for this student
    res_fe = await db.execute(
        select(TutorialFeedback)
        .filter(TutorialFeedback.user_id == sid)
        .order_by(TutorialFeedback.created_at.desc())
    )
    feedback_entries = res_fe.scalars().all()

    total = len(feedback_entries)

    if total == 0:
        return {
            "student": student_info,
            "totalEntries": 0,
            "summary": {
                "mostCommonReason": None,
                "mostProblematicStep": None,
            },
            "reasonBreakdown": [],
            "stepBreakdown": [],
            "entries": [],
            "tutorialGroups": [],
        }

    # Aggregate reasons and step types
    reason_counts: dict[str, int] = {}
    step_counts: dict[str, int] = {}

    for fb in feedback_entries:
        step_counts[fb.step_type] = step_counts.get(fb.step_type, 0) + 1
        for reason in (fb.reasons or []):
            reason_counts[reason] = reason_counts.get(reason, 0) + 1

    # Sort by frequency
    reason_breakdown = sorted(
        [{"key": k, "count": v, "percentage": round((v / total) * 100) if total > 0 else 0}
         for k, v in reason_counts.items()],
        key=lambda x: x["count"],
        reverse=True,
    )

    step_breakdown = sorted(
        [{"stepType": k, "count": v} for k, v in step_counts.items()],
        key=lambda x: x["count"],
        reverse=True,
    )

    most_common_reason = reason_breakdown[0]["key"] if reason_breakdown else None
    most_problematic_step = step_breakdown[0]["stepType"] if step_breakdown else None

    # Build entries with tutorial info
    entries = []
    # Group by tutorial for the grouped view
    tutorial_groups: dict[str, dict] = {}

    tutorial_ids = list({fb.tutorial_id for fb in feedback_entries})
    if tutorial_ids:
        res_tuts = await db.execute(select(Tutorial.id, Tutorial.title, Tutorial.tutorial_number).filter(Tutorial.id.in_(tutorial_ids)))
        tut_map = {t.id: {"title": t.title, "number": t.tutorial_number} for t in res_tuts.all()}
    else:
        tut_map = {}

    for fb in feedback_entries:
        tm = tut_map.get(fb.tutorial_id, {"title": "Unknown", "number": 0})
        
        entry = {
            "id": str(fb.id),
            "tutorialId": str(fb.tutorial_id),
            "tutorialNumber": tm["number"],
            "tutorialTitle": tm["title"],
            "stepType": fb.step_type,
            "reasons": fb.reasons or [],
            "otherText": fb.other_text,
            "createdAt": fb.created_at.isoformat() if fb.created_at else None,
        }
        entries.append(entry)

        # Group by tutorial
        tid = str(fb.tutorial_id)
        if tid not in tutorial_groups:
            tutorial_groups[tid] = {
                "tutorialId": tid,
                "tutorialNumber": entry["tutorialNumber"],
                "tutorialTitle": entry["tutorialTitle"],
                "issueCount": 0,
                "entries": [],
            }
        tutorial_groups[tid]["issueCount"] += 1
        tutorial_groups[tid]["entries"].append(entry)

    # Sort tutorial groups by issue count (most problematic first)
    sorted_groups = sorted(
        tutorial_groups.values(),
        key=lambda g: g["issueCount"],
        reverse=True,
    )

    return {
        "student": student_info,
        "totalEntries": total,
        "summary": {
            "mostCommonReason": most_common_reason,
            "mostProblematicStep": most_problematic_step,
        },
        "reasonBreakdown": reason_breakdown,
        "stepBreakdown": step_breakdown,
        "entries": entries,
        "tutorialGroups": sorted_groups,
    }


# ============================================================
# Comparison Metrics Route
# ============================================================

@router.get("/comparison-metrics")
async def get_comparison_metrics(
    state: Optional[str] = Query(None),
    use_mock: Optional[bool] = Query(None),
    user: User = Depends(require_admin),
    db: AsyncSession = Depends(get_db),
):
    """
    Returns aggregated metrics for programmes, courses, and batches.
    Filters by state if provided. Falls back to deterministic mock data in Demo Mode
    (explicit use_mock, or an empty database) so the comparison chart is never blank.
    """
    # Cache the real computed result (skip caching when explicitly forcing mock data).
    cache_key = f"admin:comparison:{state or 'all'}"
    if not use_mock:
        cached = await get_cached(cache_key)
        if cached is not None:
            return cached

    # 1. Fetch active programmes
    prog_res = await db.execute(select(Programme).filter(Programme.is_active == True).order_by(Programme.display_order))
    programmes = prog_res.scalars().all()

    if not programmes:
        # Fallback to static mock data if DB has no programmes seeded
        return get_mock_comparison_data(state)

    # 2. Fetch all courses
    course_res = await db.execute(select(Course))
    courses = course_res.scalars().all()
    course_name_map = {c.id: c.title for c in courses}

    # 3. Fetch all batches
    batch_res = await db.execute(select(Batch))
    batches = batch_res.scalars().all()

    # 4. Get many-to-many links between programmes and courses
    pc_res = await db.execute(select(ProgrammeCourse))
    programme_courses = pc_res.scalars().all()

    # Group course IDs by programme ID (from the M2M table, plus any implied by batches)
    prog_course_ids: dict = {}
    for pc in programme_courses:
        prog_course_ids.setdefault(pc.programme_id, [])
        if pc.course_id not in prog_course_ids[pc.programme_id]:
            prog_course_ids[pc.programme_id].append(pc.course_id)
    for b in batches:
        if b.programme_id:
            prog_course_ids.setdefault(b.programme_id, [])
            if b.course_id not in prog_course_ids[b.programme_id]:
                prog_course_ids[b.programme_id].append(b.course_id)

    # Determine if we should run in demo/mock mode
    tot_students_res = await db.execute(select(func.count(UserCourse.user_id)))
    tot_students = tot_students_res.scalar() or 0
    is_demo = bool(use_mock or tot_students == 0)

    # Colors for programs
    colors = {
        "fdp": "hsl(262, 83%, 58%)",
        "swayam_plus": "hsl(221, 83%, 53%)",
        "ilw": "hsl(142, 76%, 36%)"
    }

    # Per-course tutorial counts in ONE query (was a COUNT per batch).
    tut_counts = {
        cid: cnt for cid, cnt in (await db.execute(
            select(Tutorial.course_id, func.count(Tutorial.id)).group_by(Tutorial.course_id)
        )).all()
    }

    # Every student across every active-programme batch in ONE grouped query, carrying the
    # per-student COURSE-SCOPED xp (only this batch's course's tutorial XP, matching every other
    # dashboard) and the completed count (was 2 queries PER batch).
    active_prog_ids = {p.id for p in programmes}
    relevant_batch_ids = [b.id for b in batches if b.programme_id in active_prog_ids]
    rows_by_batch: dict = {}
    if relevant_batch_ids:
        merged_q = (
            select(
                UserCourse.batch_id.label("batch_id"),
                User.id.label("uid"),
                User.college_id.label("college_id"),
                course_xp_sum().label("xp"),
                func.count(case((UserTutorial.practiced == True, Tutorial.id), else_=None)).label("completed"),
            )
            .select_from(User)
            .join(UserCourse, UserCourse.user_id == User.id)
            .join(Batch, Batch.id == UserCourse.batch_id)
            .outerjoin(UserTutorial, UserTutorial.user_id == User.id)
            .outerjoin(Tutorial, (Tutorial.id == UserTutorial.tutorial_id) & (Tutorial.course_id == Batch.course_id))
            .filter(
                UserCourse.batch_id.in_(relevant_batch_ids),
                User.role == "student",
                User.is_test_account != True,
            )
            .group_by(UserCourse.batch_id, User.id, User.college_id)
        )
        if state:
            merged_q = merged_q.join(College, College.id == User.college_id).filter(
                College.location.ilike(f"%{state}%")
            )
        for row in (await db.execute(merged_q)).all():
            rows_by_batch.setdefault(row.batch_id, []).append(row)

    programmes_list = []

    for prog in programmes:
        c_ids = prog_course_ids.get(prog.id, [])
        courses_list = []

        # Aggregated metrics for this programme (weighted so real + demo batches combine).
        prog_total_colleges = 0
        prog_total_students = 0
        prog_weighted_progress_sum = 0
        prog_weighted_xp_sum = 0
        prog_batch_count = 0

        for course_id in c_ids:
            course_batches = [b for b in batches if b.course_id == course_id and b.programme_id == prog.id]
            batches_list = []

            # A course linked to a programme but with no batches gets a demo batch in Demo Mode.
            if not course_batches and is_demo:
                mock_batch_id = f"{course_id}-mock-batch"
                mock_metrics = generate_deterministic_mock_metrics(prog.id, course_id, mock_batch_id, state)
                batches_list.append({
                    "batchId": mock_batch_id,
                    "batchName": "Batch 1 (Demo)",
                    "metrics": mock_metrics
                })
                prog_total_colleges += mock_metrics["colleges"]
                prog_total_students += mock_metrics["students"]
                prog_weighted_progress_sum += mock_metrics["progress"] * mock_metrics["students"]
                prog_weighted_xp_sum += mock_metrics["avgXP"] * mock_metrics["students"]
                prog_batch_count += 1

            for batch in course_batches:
                # Look up this batch's pre-fetched student rows (was 2 queries per batch).
                batch_rows = rows_by_batch.get(batch.id, [])
                metrics = None

                if batch_rows:
                    total_tutorials = tut_counts.get(batch.course_id, 0) or 1
                    batch_student_ids = {r.uid for r in batch_rows}
                    batch_college_ids = {r.college_id for r in batch_rows if r.college_id}
                    batch_xp_sum = sum(r.xp for r in batch_rows)
                    batch_avg_xp = round(batch_xp_sum / len(batch_rows))
                    batch_progress_sum = sum(round((r.completed / total_tutorials) * 100) for r in batch_rows)
                    batch_avg_progress = round(batch_progress_sum / len(batch_rows))
                    metrics = {
                        "colleges": len(batch_college_ids),
                        "students": len(batch_student_ids),
                        "progress": batch_avg_progress,
                        "avgXP": batch_avg_xp
                    }
                else:
                    # No real students in this batch: demo values in Demo Mode, else honest zeros.
                    if is_demo:
                        metrics = generate_deterministic_mock_metrics(prog.id, course_id, batch.id, state)
                    else:
                        metrics = {"colleges": 0, "students": 0, "progress": 0, "avgXP": 0}

                if metrics:
                    batches_list.append({
                        "batchId": batch.id,
                        "batchName": batch.name,
                        "metrics": metrics
                    })
                    prog_total_colleges += metrics["colleges"]
                    prog_total_students += metrics["students"]
                    prog_weighted_progress_sum += metrics["progress"] * metrics["students"]
                    prog_weighted_xp_sum += metrics["avgXP"] * metrics["students"]
                    prog_batch_count += 1

            if batches_list:
                courses_list.append({
                    "courseId": course_id,
                    "courseName": course_name_map.get(course_id, course_id.replace("swayam_plus-", "").replace("ilw-", "").capitalize()),
                    "batches": batches_list
                })

        # Aggregate programme metrics (weighted by students so demo + real combine sensibly)
        if prog_total_students > 0:
            prog_avg_progress = round(prog_weighted_progress_sum / prog_total_students)
            prog_avg_xp = round(prog_weighted_xp_sum / prog_total_students)
        else:
            prog_avg_progress = 0
            prog_avg_xp = 0

        programmes_list.append({
            "id": prog.id,
            "name": prog.name,
            "color": colors.get(prog.id, prog.color or "hsl(200, 80%, 50%)"),
            "metrics": {
                "colleges": prog_total_colleges,
                "students": prog_total_students,
                "progress": prog_avg_progress,
                "avgXP": prog_avg_xp
            },
            "courses": courses_list
        })

    result = {"programmes": programmes_list}
    if not use_mock:
        await set_cached(cache_key, result, ttl=300)
    return result

