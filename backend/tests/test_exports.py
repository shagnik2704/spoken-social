"""
Guardrail tests for the three .xlsx export routes, which all render through
excel_export.write_styled_sheet / workbook_response.

These assert the bytes on the wire actually parse as a workbook and carry the expected
sheet names and header row. The unit-level styling equivalence was checked when the helper
was extracted; what these add is that the routes are still wired to it correctly and still
return a real file rather than, say, a JSON error body with a 200.
"""

import io
import uuid

import openpyxl
import pytest

from models.college import College
from models.course import Course
from models.user import User
from services.jwt import create_access_token

pytestmark = pytest.mark.asyncio(loop_scope="session")

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EXPORTS = [
    (
        "/api/coordinator/colleges/c1/export",
        ["Student Progress"],
        ["S.No.", "Student Name", "Progress Percentage", "Tutorials Completed", "Mini-Tasks Opted"],
    ),
    (
        "/api/admin/colleges/export",
        ["Colleges Report"],
        ["S.No.", "College Name", "Total Students Enrolled", "Active + Completed",
         "Fully Completed Students", "Students Opted for Minitasks", "Average Progress Percentage"],
    ),
    (
        "/api/admin/minitask-students/export",
        ["Mini-Task Opt-ins", "Unique Students"],
        ["S.No.", "Student Name", "Email", "College", "Tutorial #", "Tutorial Title"],
    ),
]


@pytest.fixture
async def admin_headers(db):
    db.add_all([College(id="c1", name="College One"), Course(id="linux", title="Linux")])
    await db.flush()
    admin = User(
        id=uuid.uuid4(),
        username=f"@{uuid.uuid4().hex[:10]}",
        email=f"{uuid.uuid4().hex[:10]}@example.org",
        display_name="Admin",
        role="admin",
    )
    db.add(admin)
    await db.commit()
    token = create_access_token(str(admin.id), effective_role="admin")
    return {"Authorization": f"Bearer {token}"}


@pytest.mark.parametrize("url,expected_sheets,expected_header", EXPORTS)
async def test_export_returns_a_parseable_workbook(client, admin_headers, url, expected_sheets, expected_header):
    resp = await client.get(url, headers=admin_headers)
    assert resp.status_code == 200, f"{url} returned {resp.status_code}"
    assert resp.headers["content-type"].startswith(XLSX_MEDIA_TYPE)
    assert ".xlsx" in resp.headers["content-disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(resp.content))
    assert wb.sheetnames == expected_sheets

    sheet = wb[expected_sheets[0]]
    header = [c.value for c in next(sheet.iter_rows(min_row=1, max_row=1))]
    assert header == expected_header


@pytest.mark.parametrize("url,expected_sheets,_header", EXPORTS)
async def test_export_header_row_keeps_its_styling(client, admin_headers, url, expected_sheets, _header):
    """The header styling is the part most likely to be silently lost in a refactor."""
    resp = await client.get(url, headers=admin_headers)
    sheet = openpyxl.load_workbook(io.BytesIO(resp.content))[expected_sheets[0]]
    cell = sheet.cell(row=1, column=1)
    assert cell.font.bold is True
    assert cell.font.color.rgb.endswith("FFFFFF")
    assert cell.fill.start_color.rgb.endswith("4472C4")
    assert cell.alignment.horizontal == "center"
    assert sheet.row_dimensions[1].height == 28
